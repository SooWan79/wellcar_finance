"""매출/지출 엑셀(.xlsx) 내보내기와 엑셀 입력 양식 (openpyxl).

내역 시트의 열 이름은 엑셀 올리기가 읽는 이름과 같습니다. 그래서 내보낸 파일을
이 시스템에 다시 올려도 중복 없이 같은 결과가 되고(왕복 호환), '관리번호' 열로
엑셀에서 고친 줄을 찾아 그 내역에 반영합니다.
"""
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

WEEKDAYS = ["월", "화", "수", "목", "금", "토", "일"]

FMT_WON = "#,##0"
FMT_DEC = "#,##0.00"
FMT_DATE = "yyyy-mm-dd"
FMT_PCT = "0.0%"

# (헤더, 필드, 형식, 열 너비)
INCOME_COLUMNS = [
    ("영업일자", "trx_date", "date", 12),
    ("요일", "weekday", "text", 6),
    ("매출유형", "income_type", "text", 11),
    ("매출구분", "category", "text", 14),
    ("브랜드", "manufacturer", "text", 11),
    ("차종", "car_model", "text", 13),
    ("제품모델", "product_model", "text", 14),
    ("거래처명", "client", "text", 20),
    ("거래통화", "currency", "text", 9),
    ("기준환율", "exchange_rate", "rate", 10),
    ("수량", "quantity", "qty", 7),
    ("단가", "unit_price", "price", 12),
    ("매출액", "amount", "won", 13),
    ("부가세", "vat", "won", 11),
    ("순매출액", "net_amount", "won", 13),
    ("결제유형", "payment_type", "text", 10),
    ("세금계산서발행유무", "tax_invoice", "text", 11),
    ("계좌", "account", "text", 20),
    ("미수금", "receivable", "won", 12),
    ("미지급금", "payout", "won", 12),
    ("미지급 잔액", "payout_due", "won", 12),
    ("적요", "memo", "text", 36),
    ("관리번호", "id", "id", 9),
]
EXPENSE_COLUMNS = [
    ("영업일자", "trx_date", "date", 12),
    ("요일", "weekday", "text", 6),
    ("지출유형", "expense_type", "text", 11),
    ("거래품목", "item", "text", 14),
    ("결제유형", "payment_type", "text", 10),
    ("거래처명", "client", "text", 20),
    ("거래통화", "currency", "text", 9),
    ("기준환율", "exchange_rate", "rate", 10),
    ("수량", "quantity", "qty", 7),
    ("단가", "unit_price", "price", 12),
    ("지출금액", "amount", "won", 13),
    ("미지급금", "payable", "won", 12),
    ("적요", "memo", "text", 36),
    ("관리번호", "id", "id", 9),
]

_HEAD_FONT = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="2A78D6")
_HEAD_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)
_MUTED = Font(color="5B6675")
_TOTAL_FILL = PatternFill("solid", fgColor="EEF1F6")
_TOP_LINE = Border(top=Side(style="thin", color="8B95A4"))
_WRAP = Alignment(wrap_text=True, vertical="top")


def _cell(ws, value, fmt=None, font=None, fill=None, align=None, border=None):
    c = WriteOnlyCell(ws, value=value)
    if isinstance(value, str) and value.startswith("="):
        c.data_type = "s"  # '='로 시작하는 적요 등이 수식으로 실행되지 않게 글자로 고정
    if fmt:
        c.number_format = fmt
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def _head_row(ws, titles):
    return [_cell(ws, t, font=_HEAD_FONT, fill=_HEAD_FILL, align=_HEAD_ALIGN) for t in titles]


def _as_date(s):
    try:
        return datetime.strptime(str(s), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return s


def _num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _whole(v):
    """정수로 떨어지는 값은 int로(엑셀에서 1.0이 아니라 1로 보이게)."""
    f = _num(v)
    return int(f) if f.is_integer() else f


def _value_cell(ws, kind, field, row):
    v = row[field] if field != "weekday" else None
    if kind == "date":
        return _cell(ws, _as_date(v), FMT_DATE)
    if field == "weekday":
        d = _as_date(row["trx_date"])
        return _cell(ws, WEEKDAYS[d.weekday()] if isinstance(d, date) else "")
    if kind == "won":
        return _cell(ws, _whole(v), FMT_WON)
    if kind == "id":
        return _cell(ws, int(v))
    if kind == "rate":
        return _cell(ws, _whole(v) or 1, FMT_DEC)
    if kind == "qty":
        q = _whole(v)
        return _cell(ws, q, FMT_WON if isinstance(q, int) else "#,##0.###")
    if kind == "price":
        p = _whole(v)
        krw = (row["currency"] or "KRW") == "KRW"
        return _cell(ws, p, FMT_WON if krw and isinstance(p, int) else FMT_DEC)
    return _cell(ws, "" if v is None else str(v))


def _detail_sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title)
    for i, (_h, _f, _k, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "B2"
    ws.append(_head_row(ws, [c[0] for c in columns]))
    n = 0
    for row in rows:
        ws.append([_value_cell(ws, kind, field, row) for _h, field, kind, _w in columns])
        n += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{n + 1}"
    return n


def _summary_sheet(wb, meta, totals, include_inc, include_exp):
    ws = wb.create_sheet("요약")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 10
    ws.append([_cell(ws, meta["title"], font=_TITLE)])
    # '시스템 식별자'·'내보낸 시각'은 다시 올릴 때 관리번호가 이 시스템 것인지, 그 뒤에 시스템에서
    # 고친 내역인지 알아보는 데 쓴다 (엑셀 올리기가 이 이름을 찾아 읽음)
    for label, key in (("조회 기간", "period"), ("조건", "conditions"),
                       ("내보낸 시각", "exported_at"), ("내보낸 사람", "exported_by"),
                       ("시스템 식별자", "instance")):
        if meta.get(key):
            ws.append([_cell(ws, label, font=_MUTED), _cell(ws, meta[key])])
    ws.append([])
    ws.append(_head_row(ws, ["항목", "금액(원)", "건수"]))
    lines = []
    if include_inc:
        lines += [("매출액 합계 (총매출)", totals["income"], totals["income_count"]),
                  ("미지급금 (매출차감)", totals["payout"], None),
                  ("실매출 (매출액 − 미지급금)", totals["income"] - totals["payout"], None),
                  ("부가세", totals["vat"], None),
                  ("순매출액 (공급가액)", totals["net_income"], None),
                  ("미수금 잔액 (받을 돈)", totals["receivable"], totals["receivable_count"]),
                  ("거래처 미지급 잔액 (줄 돈)", totals["payout_due"], totals["payout_due_count"])]
    if include_exp:
        lines += [("지출금액 합계", totals["expense"], totals["expense_count"]),
                  ("지출 미지급금 잔액 (줄 돈)", totals["payable"], totals["payable_count"])]
    if include_inc and include_exp:
        lines.append(("영업이익 (실매출 − 지출)", totals["income"] - totals["payout"] - totals["expense"], None))
    for label, amount, count in lines:
        ws.append([_cell(ws, label, font=_BOLD), _cell(ws, _whole(amount), FMT_WON),
                   _cell(ws, count, FMT_WON) if count is not None else _cell(ws, None)])
    ws.append([])
    for note in ("※ '매출내역'·'지출내역' 시트는 이 시스템의 [데이터 관리 → 엑셀로 입력하기]로 다시 올릴 수 있습니다.",
                 "   이미 있는 내역은 건너뛰고, 엑셀에서 고친 줄은 '관리번호'로 찾아 그 내역을 고칩니다.",
                 "   관리번호 칸은 고치거나 지우지 마세요. 새로 적는 줄은 관리번호를 비워 두면 새 내역으로 등록됩니다."):
        ws.append([_cell(ws, note, font=_MUTED)])


def _monthly_sheet(wb, monthly, include_inc, include_exp):
    ws = wb.create_sheet("월별손익")
    # '영업월'·'매출액'·'지출금액' 머리글: 다시 올릴 때 이 표를 월별 합계로 인식해 내역과 자동 대조한다
    cols = [("영업월", "month", None, 10)]
    if include_inc:
        cols += [("매출 건수", "income_count", FMT_WON, 9), ("매출액", "income", FMT_WON, 14),
                 ("매출차감(미지급금)", "payout", FMT_WON, 13), ("실매출", "sales", FMT_WON, 14),
                 ("부가세", "vat", FMT_WON, 12), ("순매출액", "net_income", FMT_WON, 14),
                 ("미수금", "receivable", FMT_WON, 12)]
    if include_exp:
        cols += [("지출 건수", "expense_count", FMT_WON, 9), ("지출금액", "expense", FMT_WON, 14),
                 ("지출 미지급금", "payable", FMT_WON, 12)]
    if include_inc and include_exp:
        cols += [("영업이익", "profit", FMT_WON, 14)]
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = c[3]
    ws.freeze_panes = "B2"
    ws.append(_head_row(ws, [c[0] for c in cols]))
    total = {}
    for m in monthly:
        ws.append([_cell(ws, m[key], fmt) if fmt else _cell(ws, m[key]) for _t, key, fmt, _w in cols])
        for _t, key, fmt, _w in cols[1:]:
            total[key] = total.get(key, 0) + (m[key] or 0)
    if monthly:
        row = [_cell(ws, "합계", font=_BOLD, fill=_TOTAL_FILL, border=_TOP_LINE)]
        row += [_cell(ws, total[key], fmt, font=_BOLD, fill=_TOTAL_FILL, border=_TOP_LINE)
                for _t, key, fmt, _w in cols[1:]]
        ws.append(row)


def _dimension_sheet(wb, spec):
    """분류별 집계표. spec: {title, label, note, cols: [키], col_labels: [이름],
    rows: [{name, cnt, total, cells: {키: 금액}}]} — 금액 큰 순으로 이미 정렬되어 있음."""
    ws = wb.create_sheet(spec["title"])
    cols = spec["cols"]
    widths = [28, 8, 15, 8] + [13] * len(cols)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.append([_cell(ws, spec["title"], font=_TITLE)])
    ws.append([_cell(ws, spec["note"], font=_MUTED)])
    ws.append(_head_row(ws, [spec["label"], "건수", "합계", "비중"] + spec["col_labels"]))
    ws.freeze_panes = "E4"
    grand = sum(r["total"] for r in spec["rows"]) or 0
    for r in spec["rows"]:
        share = (r["total"] / grand) if grand else 0
        ws.append([_cell(ws, r["name"]), _cell(ws, r["cnt"], FMT_WON), _cell(ws, _whole(r["total"]), FMT_WON),
                   _cell(ws, share, FMT_PCT)] + [_cell(ws, _whole(r["cells"].get(k, 0)), FMT_WON) for k in cols])
    if spec["rows"]:
        col_tot = {k: sum(r["cells"].get(k, 0) for r in spec["rows"]) for k in cols}
        style = {"font": _BOLD, "fill": _TOTAL_FILL, "border": _TOP_LINE}
        ws.append([_cell(ws, "합계", **style), _cell(ws, sum(r["cnt"] for r in spec["rows"]), FMT_WON, **style),
                   _cell(ws, _whole(grand), FMT_WON, **style), _cell(ws, 1 if grand else 0, FMT_PCT, **style)]
                  + [_cell(ws, _whole(col_tot[k]), FMT_WON, **style) for k in cols])


def build_workbook(meta, totals, monthly, incomes=None, expenses=None, dims=()):
    """요약·월별손익·분류별 집계·매출내역·지출내역 시트로 된 xlsx 바이트를 만든다.

    incomes/expenses는 행(dict 또는 sqlite3.Row)의 반복자이며, None이면 그 시트를 만들지 않는다.
    dims는 _dimension_sheet에 넘길 분류별 집계표 목록.
    """
    include_inc, include_exp = incomes is not None, expenses is not None
    wb = Workbook(write_only=True)
    _summary_sheet(wb, meta, totals, include_inc, include_exp)
    _monthly_sheet(wb, monthly, include_inc, include_exp)
    for spec in dims:
        _dimension_sheet(wb, spec)
    if include_inc:
        _detail_sheet(wb, "매출내역", INCOME_COLUMNS, incomes)
    if include_exp:
        _detail_sheet(wb, "지출내역", EXPENSE_COLUMNS, expenses)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- 엑셀 입력 양식
# (머리글, 열 너비, 드롭다운 목록 키, 표시 형식) — 머리글은 엑셀 올리기가 읽는 이름과 같다
TEMPLATE_INCOME = [
    ("영업일자", 12, None, FMT_DATE),
    ("매출유형", 11, "income_type", None),
    ("매출구분", 14, "income_category", None),
    ("브랜드", 11, "manufacturer", None),
    ("차종", 14, "car_model", None),
    ("제품모델", 12, None, None),
    ("거래처명", 20, "client", None),
    ("매출액", 13, None, FMT_WON),
    ("결제유형", 10, "payment_type", None),
    ("세금계산서발행유무", 11, "yn", None),
    ("계좌", 20, "account", None),
    ("미수금", 11, None, FMT_WON),
    ("미지급금", 11, None, FMT_WON),
    ("미지급금 지급완료", 11, "yn", None),
    ("적요", 32, None, None),
    ("부가세", 11, None, FMT_WON),
    ("거래통화", 9, "currency", None),
    ("CNY 기준환율", 11, None, FMT_DEC),
    ("수량", 7, None, None),
    ("단가", 11, None, None),
]
TEMPLATE_EXPENSE = [
    ("영업일자", 12, None, FMT_DATE),
    ("지출유형", 11, "expense_type", None),
    ("거래품목", 14, "expense_item", None),
    ("결제유형", 10, "payment_type", None),
    ("거래처명", 20, "client", None),
    ("지출금액", 13, None, FMT_WON),
    ("미지급금", 11, None, FMT_WON),
    ("적요", 32, None, None),
    ("거래통화", 9, "currency", None),
    ("CNY 기준환율", 11, None, FMT_DEC),
    ("수량", 7, None, None),
    ("단가", 11, None, None),
]
TEMPLATE_CODE_LISTS = [
    ("income_type", "매출유형"), ("income_category", "매출구분"), ("manufacturer", "브랜드"),
    ("car_model", "차종"), ("client", "거래처명"), ("payment_type", "결제유형"), ("account", "계좌"),
    ("currency", "거래통화"), ("expense_type", "지출유형"), ("expense_item", "거래품목"), ("yn", "Y/N"),
]
TEMPLATE_ROWS = 1000
TEMPLATE_GUIDE = [
    ("웰카오디오 매출/지출 입력 양식", "title"),
    ("", None),
    ("쓰는 방법", "head"),
    ("1. '매출입력'·'지출입력' 시트에 한 줄에 한 건씩 적습니다. 첫 줄(머리글)은 지우거나 바꾸지 마세요.", None),
    ("2. 머리글 아래 칸을 누르면 목록(드롭다운)이 나오는 칸은 목록에서 고르면 됩니다. "
     "목록에 없는 값도 적을 수 있고, 올리면 코드관리에 자동으로 추가됩니다.", None),
    ("3. 다 적었으면 시스템의 [데이터 관리 → 엑셀로 입력하기]에 이 파일을 올립니다. 새 내역은 바로 반영되고, "
     "이미 올린 줄은 건너뜁니다.", None),
    ("4. 같은 파일에 계속 이어 적고 다시 올려도 됩니다. 이미 올린 줄을 엑셀에서 고치거나 지웠다면 "
     "'엑셀 장부와 맞추기'(관리자)로 올리면 시스템도 똑같이 바뀝니다.", None),
    ("", None),
    ("칸 설명", "head"),
    ("영업일자: 2026-10-04 처럼 적습니다.", None),
    ("매출액·지출금액: 원화 금액(부가세 포함)입니다.", None),
    ("브랜드·차종: 차량 브랜드(차량제조사)와 차종입니다. 매출 집계를 브랜드별·차종별로 볼 수 있습니다.", None),
    ("매출구분: 서비스 구분(앰프수리, 탈부착 등)입니다.", None),
    ("부가세: 비워 두면 결제유형이 '카드'이거나 세금계산서가 'Y'일 때 매출액 × 10/110으로 자동 계산합니다.", None),
    ("미수금: 매출 중 아직 받지 못한 금액입니다.", None),
    ("미지급금(매출): 이 매출에서 거래처에 줘야 할 금액입니다. 실매출(매출액 − 미지급금)에서 빠집니다. "
     "이미 줬으면 '미지급금 지급완료'에 Y를 적습니다.", None),
    ("미지급금(지출): 외상으로 사서 아직 지급하지 않은 금액입니다.", None),
    ("거래통화가 CNY이면 수량·단가·CNY 기준환율을 적고, 매출액(지출금액)에는 원화로 바꾼 금액을 적습니다.", None),
]


def build_template(code_lists):
    """엑셀 입력 양식(xlsx 바이트). code_lists: {목록 키: [값, …]} — 드롭다운은 '코드목록' 시트를 가리킨다."""
    wb = Workbook()
    guide = wb.active
    guide.title = "작성 안내"
    guide.column_dimensions["A"].width = 120
    for i, (text, kind) in enumerate(TEMPLATE_GUIDE, start=1):
        c = guide.cell(row=i, column=1, value=text)
        c.font = _TITLE if kind == "title" else _BOLD if kind == "head" else Font()
        c.alignment = _WRAP
    inc = wb.create_sheet("매출입력")
    exp = wb.create_sheet("지출입력")
    codes = wb.create_sheet("코드목록")
    where = {}
    for i, (key, label) in enumerate(TEMPLATE_CODE_LISTS, start=1):
        letter = get_column_letter(i)
        head = codes.cell(row=1, column=i, value=label)
        head.font, head.fill, head.alignment = _HEAD_FONT, _HEAD_FILL, _HEAD_ALIGN
        codes.column_dimensions[letter].width = 16
        values = code_lists.get(key) or []
        for r, v in enumerate(values, start=2):
            codes.cell(row=r, column=i, value=v)
        where[key] = (letter, len(values))
    codes.freeze_panes = "A2"
    for ws, spec in ((inc, TEMPLATE_INCOME), (exp, TEMPLATE_EXPENSE)):
        for c, (head, width, code, fmt) in enumerate(spec, start=1):
            letter = get_column_letter(c)
            cell = ws.cell(row=1, column=c, value=head)
            cell.font, cell.fill, cell.alignment = _HEAD_FONT, _HEAD_FILL, _HEAD_ALIGN
            ws.column_dimensions[letter].width = width
            if fmt:
                for r in range(2, TEMPLATE_ROWS + 2):
                    ws.cell(row=r, column=c).number_format = fmt
            if code and where.get(code, (None, 0))[1]:
                src, n = where[code]
                # 목록 밖의 값도 적을 수 있게 오류 창은 띄우지 않는다 (새 거래처·차종 등)
                dv = DataValidation(type="list", formula1=f"'코드목록'!${src}$2:${src}${n + 1}",
                                    allow_blank=True, showErrorMessage=False)
                dv.add(f"{letter}2:{letter}{TEMPLATE_ROWS + 1}")
                ws.add_data_validation(dv)
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 30
    wb.active = 1  # 열면 '매출입력' 시트부터 보이게
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
