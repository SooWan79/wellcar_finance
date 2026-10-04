"""백엔드 자동 테스트 (표준 라이브러리 unittest).

    .venv/bin/python -m unittest discover -s tests -v

임시 폴더의 DB를 쓰므로 실데이터(wellcar.db)를 건드리지 않습니다.
"""
import io
import os
import sqlite3
import sys
import tempfile
import unittest

TMP = tempfile.mkdtemp(prefix="wellcar-test-")
os.environ["WELLCAR_DB"] = os.path.join(TMP, "test.db")
os.environ.pop("WELLCAR_ADMIN_PASSWORD", None)
os.environ.pop("WELLCAR_SECRET_KEY", None)
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app as wellcar  # noqa: E402  (환경변수를 먼저 정해야 함)
import backup  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

HDR = {"X-Requested-With": "XMLHttpRequest"}
ADMIN_PW = "admin-pass-1234"


def setup_code():
    db = sqlite3.connect(wellcar.DB_PATH)
    try:
        row = db.execute("SELECT value FROM settings WHERE key='setup_code'").fetchone()
        return row[0] if row else None
    finally:
        db.close()


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        wellcar.app.config["TESTING"] = True
        db = sqlite3.connect(wellcar.DB_PATH)
        if not db.execute("SELECT COUNT(*) FROM users").fetchone()[0]:
            db.close()
            c = wellcar.app.test_client()
            r = c.post("/api/auth/setup", headers=HDR, json={
                "setup_code": setup_code(), "username": "admin", "password": ADMIN_PW,
                "display_name": "사장"})
            assert r.status_code == 201, r.get_json()
        else:
            db.close()

    def setUp(self):
        db = sqlite3.connect(wellcar.DB_PATH)
        db.execute("DELETE FROM incomes")
        db.execute("DELETE FROM expenses")
        db.commit()
        db.close()
        wellcar._failures.clear()
        self.c = self.client("admin", ADMIN_PW)

    def client(self, username, password):
        c = wellcar.app.test_client()
        r = c.post("/api/auth/login", headers=HDR, json={"username": username, "password": password})
        self.assertEqual(r.status_code, 200, r.get_json())
        return c

    def post(self, c, path, data=None, **kw):
        return c.post(path, headers=HDR, json=data, **kw)

    def put(self, c, path, data):
        return c.put(path, headers=HDR, json=data)

    def delete(self, c, path):
        return c.delete(path, headers=HDR)

    def make_user(self, username, role):
        r = self.post(self.c, "/api/users", {"username": username, "password": "pass-word-99",
                                             "role": role, "display_name": username})
        if r.status_code == 409:
            return self.client(username, "pass-word-99")
        self.assertEqual(r.status_code, 201, r.get_json())
        return self.client(username, "pass-word-99")

    def income(self, c=None, **kw):
        data = {"trx_date": "2026-03-02", "income_type": "서비스제공", "category": "데크수리",
                "client": "개인", "amount": 110000, "vat": 10000, "net_amount": 100000,
                "payment_type": "카드", "memo": "테스트"}
        data.update(kw)
        r = self.post(c or self.c, "/api/incomes", data)
        self.assertEqual(r.status_code, 201, r.get_json())
        return r.get_json()


class AuthTest(Base):
    def test_requires_login(self):
        anon = wellcar.app.test_client()
        self.assertEqual(anon.get("/api/meta").status_code, 401)
        self.assertEqual(anon.get("/api/incomes").status_code, 401)
        st = anon.get("/api/auth/state").get_json()
        self.assertFalse(st["setup_required"])
        self.assertIsNone(st["user"])

    def test_setup_only_once(self):
        anon = wellcar.app.test_client()
        r = self.post(anon, "/api/auth/setup", {"setup_code": "XXXX-XXXX", "username": "evil",
                                                "password": "evil-pass-1"})
        self.assertEqual(r.status_code, 409)

    def test_csrf_header_required(self):
        r = self.c.post("/api/incomes", json={"trx_date": "2026-03-02", "income_type": "기타",
                                              "amount": 1000})
        self.assertEqual(r.status_code, 400)

    def test_wrong_password_and_throttle(self):
        anon = wellcar.app.test_client()
        for _ in range(5):
            r = self.post(anon, "/api/auth/login", {"username": "admin", "password": "nope-nope"})
            self.assertEqual(r.status_code, 401)
        r = self.post(anon, "/api/auth/login", {"username": "admin", "password": ADMIN_PW})
        self.assertEqual(r.status_code, 429)

    def test_logout(self):
        self.assertEqual(self.post(self.c, "/api/auth/logout").status_code, 200)
        self.assertEqual(self.c.get("/api/meta").status_code, 401)

    def test_roles(self):
        viewer = self.make_user("viewer1", "viewer")
        staff = self.make_user("staff1", "staff")
        # 조회 전용: 읽기만
        self.assertEqual(viewer.get("/api/incomes").status_code, 200)
        r = self.post(viewer, "/api/incomes", {"trx_date": "2026-03-02", "income_type": "기타",
                                               "amount": 1000})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(viewer.get("/api/export.xlsx").status_code, 200)
        # 직원: 내역 등록 가능, 관리자 기능 불가
        row = self.income(staff)
        self.assertEqual(row["created_by"], "staff1")
        self.assertEqual(staff.get("/api/users").status_code, 403)
        self.assertEqual(staff.get("/api/backups").status_code, 403)
        self.assertEqual(self.post(staff, "/api/codes", {"code_group": "client",
                                                         "code_value": "x"}).status_code, 403)
        # 엑셀 올리기: 직원은 추가·수정만, 맞추기(삭제가 있음)는 관리자만, 조회 전용은 불가
        self.assertEqual(self.post(staff, "/api/import-json", {"incomes": []}).status_code, 200)
        self.assertEqual(self.post(staff, "/api/import-json", {"incomes": [], "mode": "sync"}).status_code, 403)
        self.assertEqual(self.post(viewer, "/api/import-json", {"incomes": []}).status_code, 403)
        # 목표: 모두 보고 관리자만 수정
        self.assertEqual(viewer.get("/api/targets").status_code, 200)
        self.assertEqual(self.put(staff, "/api/targets", {"year": 2026, "months": {}}).status_code, 403)
        self.assertEqual(viewer.get("/api/template.xlsx").status_code, 200)

    def test_password_change_invalidates_other_sessions(self):
        self.make_user("pwuser", "staff")
        a = self.client("pwuser", "pass-word-99")
        b = self.client("pwuser", "pass-word-99")
        r = self.post(a, "/api/auth/password", {"current_password": "pass-word-99",
                                                "new_password": "new-pass-word-1"})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(a.get("/api/meta").status_code, 200)   # 바꾼 기기는 유지
        self.assertEqual(b.get("/api/meta").status_code, 401)   # 다른 기기는 로그아웃
        r = self.post(a, "/api/auth/password", {"current_password": "new-pass-word-1",
                                                "new_password": "pass-word-99"})
        self.assertEqual(r.status_code, 200)

    def test_last_admin_protected(self):
        users = self.c.get("/api/users").get_json()
        me = next(u for u in users if u["username"] == "admin")
        r = self.put(self.c, f"/api/users/{me['id']}", {"role": "staff"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.delete(self.c, f"/api/users/{me['id']}").status_code, 400)

    def test_deactivated_user_logged_out(self):
        s = self.make_user("temp1", "staff")
        uid = next(u["id"] for u in self.c.get("/api/users").get_json() if u["username"] == "temp1")
        self.assertEqual(self.put(self.c, f"/api/users/{uid}", {"active": False}).status_code, 200)
        self.assertEqual(s.get("/api/meta").status_code, 401)
        anon = wellcar.app.test_client()
        r = self.post(anon, "/api/auth/login", {"username": "temp1", "password": "pass-word-99"})
        self.assertEqual(r.status_code, 403)
        self.put(self.c, f"/api/users/{uid}", {"active": True})


class EntryTest(Base):
    def test_crud_and_get(self):
        row = self.income()
        r = self.c.get(f"/api/incomes/{row['id']}")
        self.assertEqual(r.get_json()["amount"], 110000)
        upd = dict(row, amount=220000, vat=20000, net_amount=200000)
        r = self.put(self.c, f"/api/incomes/{row['id']}", upd)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["updated_by"], "사장")
        self.assertEqual(self.delete(self.c, f"/api/incomes/{row['id']}").status_code, 200)
        self.assertEqual(self.c.get(f"/api/incomes/{row['id']}").status_code, 404)

    def test_pagination_and_filters(self):
        db = sqlite3.connect(wellcar.DB_PATH)
        db.executemany(
            "INSERT INTO incomes(trx_date, income_type, client, amount, vat, net_amount, memo) "
            "VALUES (?,?,?,?,0,?,?)",
            [(f"2025-{m:02d}-{d:02d}", "상품판매", f"거래처{i % 7}", 1000 + i, 1000 + i, f"메모{i}")
             for i, (m, d) in enumerate((m, d) for m in range(1, 13) for d in range(1, 29))])
        db.commit()
        db.close()
        total = 12 * 28
        r = self.c.get("/api/incomes?page=1&size=50").get_json()
        self.assertEqual(r["total"], total)
        self.assertEqual(r["pages"], 7)
        self.assertEqual(len(r["items"]), 50)
        self.assertEqual(r["items"][0]["trx_date"], "2025-12-28")
        self.assertEqual(r["sum"]["amount"], sum(1000 + i for i in range(total)))
        last = self.c.get("/api/incomes?page=7&size=50").get_json()
        self.assertEqual(len(last["items"]), total - 300)
        r = self.c.get("/api/incomes?from=2025-03-01&to=2025-03-31&size=500").get_json()
        self.assertEqual(r["total"], 28)
        r = self.c.get("/api/incomes?q=거래처3&size=500").get_json()
        self.assertEqual(r["total"], len([i for i in range(total) if i % 7 == 3]))
        r = self.c.get("/api/incomes?q=1005").get_json()  # 숫자 검색은 금액도 찾는다
        self.assertTrue(any(it["amount"] == 1005 for it in r["items"]))
        r = self.c.get("/api/incomes?q=100%25").get_json()  # LIKE 특수문자 안전
        self.assertEqual(r["total"], 0)

    def test_date_normalized_and_zero_amount_rejected(self):
        r = self.post(self.c, "/api/incomes", {"trx_date": "2026-3-5", "income_type": "기타", "amount": 1000})
        self.assertEqual(r.get_json()["trx_date"], "2026-03-05")
        r = self.post(self.c, "/api/incomes", {"trx_date": "2026-03-05", "income_type": "기타", "amount": "0"})
        self.assertEqual(r.status_code, 400)

    def test_receivable_and_settle(self):
        row = self.income(receivable=50000)
        r = self.post(self.c, "/api/incomes", {"trx_date": "2026-03-03", "income_type": "기타",
                                               "amount": 1000, "receivable": 5000})
        self.assertEqual(r.status_code, 400)  # 미수금 > 금액
        rec = self.c.get("/api/receivables").get_json()
        self.assertEqual(rec["incomes"]["total"], 50000)
        self.assertEqual(rec["incomes"]["by_client"][0]["name"], "개인")
        self.assertEqual(self.c.get("/api/incomes?outstanding=1").get_json()["total"], 1)
        r = self.post(self.c, f"/api/incomes/{row['id']}/settle")
        self.assertEqual(r.get_json()["receivable"], 0)
        self.assertEqual(self.c.get("/api/receivables").get_json()["incomes"]["total"], 0)
        summary = self.c.get("/api/receivables?summary=1").get_json()
        self.assertEqual(summary["incomes"], {"total": 0, "count": 0})

    def test_cash_counts_non_card(self):
        """현금매출 = 카드 외 전부 (화면은 amount − card로 계산)"""
        self.income(payment_type="입금", amount=30000, vat=0, net_amount=30000)
        self.income(payment_type="카드")
        r = self.c.get("/api/agg?kind=incomes&from=2026-03-02&to=2026-03-02&group=date").get_json()
        row = r["rows"][0]
        self.assertEqual((row["date"], row["card"], row["amount"] - row["card"]), ("2026-03-02", 110000, 30000))

    def test_payout_deducted_and_settled_separately(self):
        """매출의 미지급금: 거래처에 줄 돈. 지급 완료해도 매출차감 금액은 그대로."""
        row = self.income(amount=500000, vat=0, net_amount=500000, payout=150000)
        self.assertEqual((row["payout"], row["payout_due"]), (150000, 150000))  # 지급 전으로 시작
        bad = self.post(self.c, "/api/incomes", {"trx_date": "2026-03-02", "income_type": "기타",
                                                 "amount": 1000, "payout": 2000})
        self.assertEqual(bad.status_code, 400)
        bad = self.put(self.c, f"/api/incomes/{row['id']}", {"payout_due": 200000})  # 잔액 > 미지급금
        self.assertEqual(bad.status_code, 400)
        rec = self.c.get("/api/receivables").get_json()
        self.assertEqual((rec["payouts"]["total"], rec["payouts"]["count"]), (150000, 1))
        r = self.post(self.c, f"/api/incomes/{row['id']}/settle", {"field": "payout_due"}).get_json()
        self.assertEqual((r["payout"], r["payout_due"]), (150000, 0))
        self.assertEqual(self.post(self.c, f"/api/incomes/{row['id']}/settle", {"field": "amount"}).status_code, 400)
        agg = self.c.get("/api/agg?kind=incomes&from=2026-03-01&to=2026-03-31").get_json()["rows"][0]
        self.assertEqual(agg["amount"] - agg["payout"], 350000)  # 실매출
        # 일부만 보낸 수정: 저장된 값과 합쳐 검사·저장
        r = self.put(self.c, f"/api/incomes/{row['id']}", {"memo": "고침"})
        self.assertEqual((r.status_code, r.get_json()["payout"], r.get_json()["amount"]), (200, 150000, 500000))

    def test_bulk_settle(self):
        ids = [self.income(receivable=10000 * (i + 1))["id"] for i in range(3)]
        r = self.post(self.c, "/api/settle-bulk", {"table": "incomes", "field": "receivable", "ids": ids[:2]})
        self.assertEqual(r.get_json()["updated"], 2)
        self.assertEqual(self.c.get("/api/receivables").get_json()["incomes"]["total"], 30000)
        for bad in ({"table": "incomes", "field": "receivable", "ids": []},
                    {"table": "users", "ids": [1]},
                    {"table": "expenses", "field": "payout_due", "ids": [1]},
                    {"table": "incomes", "ids": ["1"]}):
            self.assertEqual(self.post(self.c, "/api/settle-bulk", bad).status_code, 400, bad)


class AggregateTargetTest(Base):
    def test_agg_groups_and_filters(self):
        self.income(trx_date="2026-01-05", manufacturer="벤츠", car_model="E클래스", amount=300000, vat=0)
        self.income(trx_date="2026-01-20", manufacturer="벤츠", car_model="S클래스", amount=200000, vat=0, payout=50000)
        self.income(trx_date="2026-02-03", manufacturer="BMW", car_model="5시리즈", amount=100000, vat=0)
        r = self.c.get("/api/agg?kind=incomes&from=2026-01-01&to=2026-12-31&group=month,manufacturer").get_json()
        got = {(x["month"], x["manufacturer"]): (x["cnt"], x["amount"], x["payout"]) for x in r["rows"]}
        self.assertEqual(got, {("2026-01", "벤츠"): (2, 500000, 50000), ("2026-02", "BMW"): (1, 100000, 0)})
        r = self.c.get("/api/agg?kind=incomes&from=2026-01-01&to=2026-12-31&group=car_model"
                       "&f.manufacturer=벤츠").get_json()
        self.assertEqual(sorted(x["car_model"] for x in r["rows"]), ["E클래스", "S클래스"])
        r = self.c.get("/api/agg?kind=incomes&from=2026-01-01&to=2026-12-31").get_json()
        self.assertEqual((r["rows"][0]["cnt"], r["rows"][0]["amount"]), (3, 600000))
        db = sqlite3.connect(wellcar.DB_PATH)
        db.execute("INSERT INTO expenses(trx_date, expense_type, item, amount, payable) "
                   "VALUES ('2026-01-07','제품원가','제품매입',40000,10000)")
        db.commit()
        db.close()
        e = self.c.get("/api/agg?kind=expenses&from=2026-01-01&to=2026-01-31&group=year").get_json()["rows"][0]
        self.assertEqual((e["year"], e["amount"], e["cost"], e["payable"]), ("2026", 40000, 40000, 10000))
        for q in ("kind=x&from=2026-01-01&to=2026-01-02", "kind=incomes&from=2026-02-01&to=2026-01-01",
                  "kind=incomes&from=2026-01-01&to=2026-01-02&group=memo",
                  "kind=incomes&from=2026-01-01&to=2026-01-02&group=date,date",
                  "kind=incomes&from=2026-01-01&to=2026-01-02&f.memo=x",
                  "kind=expenses&from=2026-01-01&to=2026-01-02&group=car_model",
                  "kind=incomes&from=0001-01-01&to=2026-01-02", "kind=incomes&from=x&to=y"):
            self.assertEqual(self.c.get(f"/api/agg?{q}").status_code, 400, q)

    def test_targets_put_get_and_validation(self):
        months = {"1": {"sales": 30000000, "profit": -500000, "expense": 20000000}, "2": {"sales": 10}}
        r = self.put(self.c, "/api/targets", {"year": 2031, "months": months})
        self.assertEqual(r.status_code, 200, r.get_json())
        t = self.c.get("/api/targets").get_json()["targets"]
        self.assertEqual(t["2031-01"], {"sales": 30000000, "profit": -500000, "expense": 20000000})
        self.assertEqual(t["2031-02"], {"sales": 10})
        # 다시 저장하면 그 해를 통째로 바꾼다 (빠진 달·빈 값은 지움)
        self.put(self.c, "/api/targets", {"year": 2031, "months": {"2": {"sales": 20, "profit": ""}}})
        t = self.c.get("/api/targets").get_json()["targets"]
        self.assertNotIn("2031-01", t)
        self.assertEqual(t["2031-02"], {"sales": 20})
        for bad in ({"year": 2031, "months": {"1": {"sales": -1}}}, {"year": "2031", "months": {}},
                    {"year": 2031, "months": {"1": {"sales": "많이"}}}, {"year": 2031, "months": {"1": 5}},
                    {"year": 2031, "months": {"1": {"expense": 1e20}}}, {"year": 99999, "months": {}}):
            self.assertEqual(self.put(self.c, "/api/targets", bad).status_code, 400, bad)
        self.put(self.c, "/api/targets", {"year": 2031, "months": {}})

    def test_seed_codes_include_car_models(self):
        codes = self.c.get("/api/codes").get_json()
        benz = [c["code_value"] for c in codes["car_model"] if c["parent_value"] == "벤츠"]
        self.assertIn("E클래스", benz)


class HardeningTest(Base):
    """이상한 입력에서도 500 오류 없이 400/403/404로 답하는지."""

    def test_bad_bodies_and_values(self):
        self.assertEqual(self.c.post("/api/incomes", headers=HDR, json=[1, 2]).status_code, 400)
        anon = wellcar.app.test_client()
        self.assertEqual(anon.post("/api/auth/login", headers=HDR, json=["admin"]).status_code, 401)
        r = self.post(self.c, "/api/incomes", {"trx_date": "2026-03-02", "income_type": "기타", "amount": 1e20})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.c.get("/api/incomes/99999999999999999999").status_code, 404)
        self.assertEqual(self.post(self.c, "/api/users", {"username": "x1", "password": "pass-word-99",
                                                          "role": ["admin"]}).status_code, 400)
        uid = self.c.get("/api/users").get_json()[0]["id"]
        self.assertEqual(self.put(self.c, f"/api/users/{uid}", {"active": "false"}).status_code, 400)
        self.assertEqual(self.post(self.c, "/api/codes", {"code_group": 1, "code_value": 2}).status_code, 201)
        for body in ({"incomes": "x"}, {"incomes": [], "mode": "replace"}, {"incomes": [1, "x", None]}):
            r = self.post(self.c, "/api/import-json", body)
            self.assertIn(r.status_code, (200, 400), body)
        self.assertEqual(self.post(self.c, "/api/incomes/1/settle", {"field": "id"}).status_code, 400)

    def test_amount_rounds_half_up_like_screen(self):
        r = self.post(self.c, "/api/incomes", {"trx_date": "2026-03-02", "income_type": "기타", "amount": 1000.5})
        self.assertEqual(r.get_json()["amount"], 1001)

    def test_setup_code_non_ascii(self):
        """한글 입력 상태로 설정 코드를 넣어도 500이 아니라 '맞지 않음'."""
        db = sqlite3.connect(wellcar.DB_PATH)
        saved = db.execute("SELECT * FROM users").fetchall()
        cols = [d[0] for d in db.execute("SELECT * FROM users").description]
        try:
            db.execute("DELETE FROM users")
            db.execute("INSERT OR REPLACE INTO settings(key, value) VALUES ('setup_code', 'ABCD-EFGH')")
            db.commit()
            anon = wellcar.app.test_client()
            r = self.post(anon, "/api/auth/setup", {"setup_code": "가나다라", "username": "owner",
                                                    "password": "owner-pass-1"})
            self.assertEqual(r.status_code, 403)
        finally:
            db.execute("DELETE FROM users")
            db.executemany(f"INSERT INTO users ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})", saved)
            db.execute("DELETE FROM settings WHERE key='setup_code'")
            db.commit()
            db.close()
            wellcar._failures.clear()


class ImportExportTest(Base):
    def rows(self):
        inc = [{"trx_date": "2026-01-02", "income_type": "서비스제공", "category": "데크수리",
                "client": "A", "amount": 750000, "vat": 0, "net_amount": 750000,
                "payment_type": "입금", "memo": "x", "receivable": 250000, "src": "매출집계!6"},
               {"trx_date": "2026-01-02", "income_type": "서비스제공", "category": "데크수리",
                "client": "A", "amount": 750000, "vat": 0, "net_amount": 750000,
                "payment_type": "입금", "memo": "x", "src": "매출집계!7"},  # 파일 안의 정당한 중복
               {"trx_date": "2026-01-05", "income_type": "상품판매", "category": "앰프", "client": "B",
                "currency": "CNY", "exchange_rate": 190.5, "quantity": 2, "unit_price": 1000,
                "amount": 381000, "vat": 34636, "net_amount": 346364, "payment_type": "카드",
                "tax_invoice": "Y", "memo": "=HYPERLINK(\"http://x\")", "src": "매출집계!8"},
               {"trx_date": "날짜아님", "income_type": "기타", "amount": 1, "src": "매출집계!9"}]
        exp = [{"trx_date": "2026-01-03", "expense_type": "생활", "item": "마트", "client": "C",
                "payment_type": "법인체크", "amount": 30400, "payable": 10000, "src": "지출집계!6"}]
        return inc, exp

    def imp(self, inc, exp=(), **kw):
        r = self.post(self.c, "/api/import-json", dict({"incomes": list(inc), "expenses": list(exp)}, **kw))
        self.assertEqual(r.status_code, 200, r.get_json())
        return r.get_json()

    def test_import_idempotent_and_fields(self):
        inc, exp = self.rows()
        r = self.imp(inc, exp)
        self.assertEqual((r["incomes"]["added"], r["incomes"]["same"], r["incomes"]["invalid"]), (3, 0, 1))
        self.assertEqual(r["expenses"]["added"], 1)
        self.assertEqual(r["samples"]["incomes"]["invalid"][0]["src"], "매출집계!9")
        r2 = self.imp(inc, exp)
        self.assertEqual((r2["incomes"]["added"], r2["incomes"]["same"]), (0, 3))
        self.assertEqual(r2["backup"], "")  # 바뀌는 게 없으면 백업도 만들지 않음
        items = self.c.get("/api/incomes?size=10").get_json()["items"]
        cny = next(i for i in items if i["currency"] == "CNY")
        self.assertEqual((cny["exchange_rate"], cny["quantity"], cny["unit_price"]), (190.5, 2, 1000))
        self.assertEqual(sum(i["receivable"] for i in items), 250000)
        self.assertTrue(all(i["source"] == "excel" and i["import_key"] for i in items))
        e = self.c.get("/api/expenses").get_json()["items"][0]
        self.assertEqual(e["payable"], 10000)
        months = self.c.get("/api/agg?kind=incomes&from=2026-01-01&to=2026-01-31&group=month").get_json()["rows"]
        self.assertEqual(months[0]["cnt"], 3)
        self.assertEqual(months[0]["amount"], 750000 * 2 + 381000)

    def test_dry_run_changes_nothing(self):
        inc, exp = self.rows()
        r = self.imp(inc, exp, dry_run=True)
        self.assertEqual((r["dry_run"], r["incomes"]["added"], r["backup"]), (True, 3, ""))
        self.assertEqual(self.c.get("/api/incomes").get_json()["total"], 0)

    def test_old_rows_without_payable_note_are_duplicates(self):
        """이전 버전이 적요에 '[미지급금 …]'을 붙였든 안 붙였든 같은 거래로 본다."""
        db = sqlite3.connect(wellcar.DB_PATH)
        db.execute("INSERT INTO incomes(trx_date, income_type, category, client, amount, vat, net_amount, memo) "
                   "VALUES ('2026-02-01','서비스제공','데크수리','A',500000,0,500000,'작업')")
        db.commit()
        db.close()
        row = {"trx_date": "2026-02-01", "income_type": "서비스제공", "category": "데크수리", "client": "A",
               "amount": 500000, "vat": 0, "net_amount": 500000, "memo": "작업 [미지급금 50,000원]"}
        r = self.imp([row])
        self.assertEqual((r["incomes"]["added"], r["incomes"]["same"]), (0, 1))

    def test_old_payout_note_migrated(self):
        """예전 가져오기가 적요에 남긴 '[미지급금 N원]'은 미지급금(매출차감) 열로 옮긴다."""
        db = sqlite3.connect(wellcar.DB_PATH)
        db.row_factory = sqlite3.Row
        db.execute("INSERT INTO incomes(trx_date, income_type, client, amount, vat, net_amount, memo, created_by) "
                   "VALUES ('2026-02-01','서비스제공','A',500000,0,500000,'작업 [미지급금 50,000원]','사장 (엑셀)')")
        wellcar._data_migrations(db)
        db.commit()
        r = db.execute("SELECT memo, payout, payout_due, source, import_key FROM incomes").fetchone()
        db.close()
        self.assertEqual((r["memo"], r["payout"], r["payout_due"], r["source"]), ("작업", 50000, 50000, "excel"))
        self.assertTrue(r["import_key"].startswith("2026-02-01"))

    def test_web_edit_of_imported_row_not_duplicated(self):
        """엑셀로 들어온 내역을 화면에서 고친 뒤 같은 파일을 다시 올려도 다시 들어오지 않는다."""
        inc, _ = self.rows()
        self.imp(inc[:1])
        row = self.c.get("/api/incomes").get_json()["items"][0]
        self.put(self.c, f"/api/incomes/{row['id']}", {"amount": 760000, "memo": "화면에서 고침"})
        r = self.imp(inc[:1])
        self.assertEqual((r["incomes"]["added"], r["incomes"]["same"]), (0, 1))

    def test_update_by_management_number(self):
        """이 시스템에서 받은 엑셀을 고쳐 올리면 관리번호로 그 내역을 고친다(충돌·삭제된 번호는 건너뜀)."""
        a = self.income(amount=100000, vat=0, net_amount=100000, memo="원래")
        b = self.income(amount=200000, vat=0, net_amount=200000, memo="둘째")
        exported = "2999-01-01 00:00:00"  # 내보낸 시각이 나중이면 충돌 아님
        rows = [dict(a, amount=150000, net_amount=150000, memo="엑셀에서 고침", ref_time=exported),
                dict(b, ref_time=exported),  # 그대로
                dict(a, id=987654, ref_time=exported)]  # 없는 관리번호
        for r in rows:
            r.pop("created_at", None)
        res = self.imp(rows)
        self.assertEqual((res["incomes"]["updated"], res["incomes"]["same"], res["incomes"]["missing"]), (1, 1, 1))
        got = self.c.get(f"/api/incomes/{a['id']}").get_json()
        self.assertEqual((got["amount"], got["memo"]), (150000, "엑셀에서 고침"))
        self.assertTrue(got["updated_by"].endswith("(엑셀)"))
        # 파일을 받은 뒤 시스템에서 고친 내역은 덮어쓰지 않는다
        stale = dict(b, amount=999000, ref_time="2000-01-01 00:00:00")
        stale.pop("created_at", None)
        res = self.imp([stale])
        self.assertEqual(res["incomes"]["conflict"], 1)
        self.assertEqual(self.c.get(f"/api/incomes/{b['id']}").get_json()["amount"], 200000)
        # 파일에 없는 열(_absent)은 건드리지 않는다
        part = {"id": b["id"], "ref_time": exported, "trx_date": b["trx_date"], "income_type": b["income_type"],
                "amount": b["amount"], "memo": "적요만", "_absent": ["client", "category", "vat", "net_amount"]}
        self.imp([part])
        got = self.c.get(f"/api/incomes/{b['id']}").get_json()
        self.assertEqual((got["memo"], got["client"], got["category"]), ("적요만", "개인", "데크수리"))

    def test_sync_mode_matches_excel_ledger(self):
        """엑셀 장부와 맞추기: 엑셀로 들어온 내역만 파일과 똑같이 (직접 입력한 내역은 그대로)."""
        inc, _ = self.rows()
        self.imp(inc[:3])
        web = self.income(trx_date="2026-01-03", amount=55000, vat=0, net_amount=55000, memo="화면 입력")
        edited = [inc[0], dict(inc[2], amount=400000, vat=36364, net_amount=363636)]  # 1줄 지움 · 1줄 금액 고침
        plan = self.imp(edited, mode="sync", dry_run=True)
        self.assertEqual((plan["incomes"]["deleted"], plan["incomes"]["added"], plan["incomes"]["same"]), (2, 1, 1))
        self.assertEqual(self.c.get("/api/incomes").get_json()["total"], 4)  # 미리보기는 그대로
        res = self.imp(edited, mode="sync")
        self.assertTrue(res["backup"].endswith("-pre-import.db"))
        items = self.c.get("/api/incomes?size=50").get_json()["items"]
        self.assertEqual(sorted(i["amount"] for i in items), [55000, 400000, 750000])
        self.assertIn(web["id"], [i["id"] for i in items])
        again = self.imp(edited, mode="sync")
        self.assertEqual((again["incomes"]["added"], again["incomes"]["deleted"], again["incomes"]["same"]), (0, 0, 2))

    def test_import_makes_pre_import_backup(self):
        self.income()
        inc, exp = self.rows()
        r = self.imp(inc, exp)
        self.assertTrue(r["backup"].endswith("-pre-import.db"))
        self.assertTrue(os.path.isfile(os.path.join(wellcar.BACKUP_DIR, r["backup"])))

    def test_export_round_trip(self):
        inc, exp = self.rows()
        self.imp(inc, exp)
        r = self.c.get("/api/export.xlsx?from=2026-01-01&to=2026-12-31")
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r.headers["Content-Disposition"])
        wb = load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb.sheetnames, ["요약", "월별손익", "브랜드별", "차종별", "서비스구분별", "지출유형별",
                                         "매출내역", "지출내역"])
        ws = wb["매출내역"]
        head = [c.value for c in ws[1]]
        for h in ("미수금", "거래통화", "브랜드", "차종", "미지급금", "미지급 잔액", "관리번호"):
            self.assertIn(h, head)
        self.assertEqual(ws.max_row, 4)  # 헤더 + 3건
        ids = sorted(ws.cell(row=i, column=head.index("관리번호") + 1).value for i in range(2, 5))
        self.assertEqual(ids, sorted(i["id"] for i in self.c.get("/api/incomes").get_json()["items"]))
        memo_col = head.index("적요")
        memos = [ws.cell(row=i, column=memo_col + 1) for i in range(2, 5)]
        formula_cell = next(c for c in memos if str(c.value).startswith("="))
        self.assertEqual(formula_cell.data_type, "s")  # 수식으로 실행되지 않음
        monthly = wb["월별손익"]
        self.assertEqual(monthly.cell(row=2, column=1).value, "2026-01")
        summary = {r[0]: r[1] for r in wb["요약"].iter_rows(values_only=True) if r and r[0]}
        meta = self.c.get("/api/meta").get_json()
        self.assertEqual(summary["시스템 식별자"], meta["instance"])
        self.assertRegex(summary["내보낸 시각"], r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")
        brands = list(wb["브랜드별"].iter_rows(values_only=True))
        self.assertEqual(brands[2][:4], ("브랜드", "건수", "합계", "비중"))
        # 조건 내보내기: 매출만, 검색어
        r = self.c.get("/api/export.xlsx?kind=incomes&q=B")
        wb = load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb.sheetnames, ["요약", "월별손익", "브랜드별", "차종별", "서비스구분별", "매출내역"])
        self.assertEqual(wb["매출내역"].max_row, 2)

    def test_template_has_dropdowns(self):
        r = self.c.get("/api/template.xlsx")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb.sheetnames, ["작성 안내", "매출입력", "지출입력", "코드목록"])
        head = [c.value for c in wb["매출입력"][1]]
        for h in ("영업일자", "매출유형", "매출액", "브랜드", "차종", "미지급금", "미지급금 지급완료"):
            self.assertIn(h, head)
        refs = [str(dv.sqref) for dv in wb["매출입력"].data_validations.dataValidation]
        self.assertIn("B2:B1001", refs)  # 매출유형 드롭다운
        codes = [c.value for c in wb["코드목록"]["A"]][1:]
        self.assertIn("서비스제공", codes)


class BackupTest(Base):
    def test_backup_restore_cycle(self):
        self.income(amount=110000)
        made = self.post(self.c, "/api/backups").get_json()
        self.assertEqual(made["kind"], "manual")
        self.income(amount=220000, vat=20000, net_amount=200000)
        self.assertEqual(self.c.get("/api/incomes").get_json()["total"], 2)
        r = self.post(self.c, f"/api/backups/{made['name']}/restore").get_json()
        self.assertEqual(r["restored"]["incomes"], 1)
        self.assertTrue(r["safety_backup"].endswith("-pre-restore.db"))
        self.assertEqual(self.c.get("/api/incomes").get_json()["total"], 1)
        # 복원해도 로그인 계정은 그대로
        self.assertEqual(self.c.get("/api/meta").status_code, 200)
        dl = self.c.get(f"/api/backups/{made['name']}/download")
        self.assertEqual(dl.status_code, 200)
        self.assertTrue(dl.data.startswith(b"SQLite format 3"))
        dl.close()
        # 업로드 → 목록에 'upload'로 보관
        up = self.c.post("/api/backups/upload", headers=HDR,
                         data={"file": (io.BytesIO(dl.data), "my.db")},
                         content_type="multipart/form-data")
        self.assertEqual(up.status_code, 201, up.get_json())
        self.assertEqual(up.get_json()["kind"], "upload")
        bad = self.c.post("/api/backups/upload", headers=HDR,
                          data={"file": (io.BytesIO(b"not a db"), "x.db")},
                          content_type="multipart/form-data")
        self.assertEqual(bad.status_code, 400)

    def test_restore_oldest_pre_restore_backup(self):
        """보관 개수가 꽉 찬 상태에서 가장 오래된 '복원 직전' 백업으로 복원해도 파일이 지워지지 않는다."""
        self.income()
        made = [backup.create_backup(wellcar.DB_PATH, wellcar.BACKUP_DIR, "pre-restore")
                for _ in range(backup.KEEP["pre-restore"])]
        oldest = [it for it in backup.list_backups(wellcar.BACKUP_DIR) if it["kind"] == "pre-restore"][-1]
        self.assertIn(oldest["name"], [m["name"] for m in made])
        r = self.post(self.c, f"/api/backups/{oldest['name']}/restore")
        self.assertEqual(r.status_code, 200, r.get_json())

    def test_path_traversal_rejected(self):
        self.assertEqual(self.c.get("/api/backups/..%2Fapp.py/download").status_code, 404)
        self.assertEqual(self.c.get("/api/backups/wellcar.db/download").status_code, 404)

    def test_restore_old_schema_backup(self):
        """열이 추가되기 전의 옛 DB 백업도 복원된다."""
        old = os.path.join(TMP, "old.db")
        db = sqlite3.connect(old)
        db.executescript("""
            CREATE TABLE incomes (id INTEGER PRIMARY KEY AUTOINCREMENT, trx_date TEXT NOT NULL,
                income_type TEXT NOT NULL, amount INTEGER NOT NULL DEFAULT 0);
            CREATE TABLE expenses (id INTEGER PRIMARY KEY AUTOINCREMENT, trx_date TEXT NOT NULL,
                expense_type TEXT NOT NULL, amount INTEGER NOT NULL DEFAULT 0);
            INSERT INTO incomes(trx_date, income_type, amount) VALUES ('2024-01-01','기타',5000);
        """)
        db.commit()
        db.close()
        counts = backup.restore_data(wellcar.DB_PATH, old)
        self.assertEqual(counts["incomes"], 1)
        it = self.c.get("/api/incomes").get_json()["items"][0]
        self.assertEqual((it["amount"], it["receivable"]), (5000, 0))

    def test_auto_backup_once_a_day(self):
        self.income()
        for it in backup.list_backups(wellcar.BACKUP_DIR):
            if it["kind"] == "auto":
                os.remove(os.path.join(wellcar.BACKUP_DIR, it["name"]))
        self.assertIsNotNone(backup.auto_backup_if_due(wellcar.DB_PATH, wellcar.BACKUP_DIR, 24))
        self.assertIsNone(backup.auto_backup_if_due(wellcar.DB_PATH, wellcar.BACKUP_DIR, 24))


if __name__ == "__main__":
    unittest.main()
