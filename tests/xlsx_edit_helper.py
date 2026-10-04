"""화면 테스트(ui_smoke.js)가 쓰는 엑셀 편집 도우미 — 사람이 엑셀에서 하는 일을 흉내 낸다.

    python tests/xlsx_edit_helper.py fill-template 양식.xlsx 결과.xlsx
        입력 양식(매출입력·지출입력)에 새 내역을 몇 줄 적는다. 부가세 칸은 비워 자동 계산을 시험한다.
    python tests/xlsx_edit_helper.py edit-export 내보낸파일.xlsx 결과.xlsx 관리번호 새금액
        내보낸 파일의 매출내역에서 그 관리번호 줄의 매출액과 적요를 고친다(관리번호는 그대로).
"""
import sys
from datetime import date

from openpyxl import load_workbook


def header_index(ws, row=1):
    return {str(c.value).strip(): c.column for c in ws[row] if c.value is not None}


def fill_template(src, dst):
    wb = load_workbook(src)
    inc, exp = wb["매출입력"], wb["지출입력"]
    hi, he = header_index(inc), header_index(exp)
    today = date.today()
    sales = [
        {"영업일자": today, "매출유형": "서비스제공", "매출구분": "앰프수리", "브랜드": "벤츠", "차종": "E클래스",
         "거래처명": "양식시험상사", "매출액": 330000, "결제유형": "카드", "적요": "양식 시험 1"},
        {"영업일자": today, "매출유형": "서비스제공", "매출구분": "탈부착", "브랜드": "BMW", "차종": "X9시험",
         "거래처명": "양식시험상사", "매출액": 500000, "결제유형": "계좌입금", "미지급금": 100000,
         "미지급금 지급완료": "N", "적요": "양식 시험 2"},
    ]
    for r, row in enumerate(sales, start=2):
        for k, v in row.items():
            inc.cell(row=r, column=hi[k], value=v)
    exp.cell(row=2, column=he["영업일자"], value=today)
    exp.cell(row=2, column=he["지출유형"], value="비용")
    exp.cell(row=2, column=he["거래품목"], value="식대")
    exp.cell(row=2, column=he["지출금액"], value=12000)
    exp.cell(row=2, column=he["적요"], value="양식 시험 지출")
    wb.save(dst)


def edit_export(src, dst, rid, amount):
    wb = load_workbook(src)
    ws = wb["매출내역"]
    h = header_index(ws)
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=h["관리번호"]).value == int(rid):
            ws.cell(row=r, column=h["매출액"], value=int(amount))
            ws.cell(row=r, column=h["적요"], value="엑셀에서 고침")
            break
    else:
        sys.exit(f"관리번호 {rid}를 찾지 못했습니다.")
    wb.save(dst)


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "fill-template":
        fill_template(sys.argv[2], sys.argv[3])
    elif cmd == "edit-export":
        edit_export(*sys.argv[2:6])
    else:
        sys.exit("알 수 없는 명령: " + cmd)
