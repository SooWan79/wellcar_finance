"""원본 경영관리 엑셀과 같은 구조의 시험용 대용량 통합문서를 만든다 (실데이터 아님).

    .venv/bin/python tests/make_sample_workbook.py --out /tmp/sample.xlsm --sales 6000 --expenses 4000

원본(2023 모두플래닛_수입지출관리.xlsm)에서 확인한 구조를 따른다.
- 매출집계·지출집계와 월별 시트(_1월~_12월): 5행이 머리글, A열은 비어 있고 6행부터 내역
- 집계 시트 2~3행의 지표 칸, 서식만 있는 빈 행(매출액·부가세·순매출액이 0)
- 매출 열: 영업년월·영업일자·요일·매출유형·매출구분·차량제조사·제품모델·거래처명·매출액·
  현금매출액·카드매출액·부가세(1/11, 반올림 안 함)·순매출액·계좌·결제유형·세금계산서발행유무·적요·미수금·미지급금
- 지출 열: 영업년월·영업일자·요일·지출유형·거래품목·결제유형·거래처명·지출금액·적요
- 'OOOO년 경영분석' 시트 15행 머리글(영업월·매출액·지출액·영업이익), 16~27행이 1~12월
- 매출분석_N월(영업일자는 있지만 매출유형이 없는 일별 표), 매출입력·지출입력(입력 양식)

--mode both: 큰 집계 시트와 월별 시트에 같은 내역이 함께 있음(시트 간 중복 제거 시험)
--mode monthly: 원본 2023 파일처럼 월별 시트에만 내역이 있음
기대값(정답)은 같은 이름의 .json 파일로 저장한다.
"""
import argparse
import json
import random
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook

WD = ["월요일", "화요일", "수요일", "목요일", "금요일", "토요일", "일요일"]
SALES_HEAD = ["영업년월", "영업일자", "요일", "매출유형", "매출구분", "차량제조사", "제품모델", "거래처명",
              "매출액", "현금매출액", "카드매출액", "부가세", "순매출액", "계좌", "결제유형",
              "세금계산서발행유무", "적요", "미수금", "미지급금"]
EXP_HEAD = ["영업년월", "영업일자", "요일", "지출유형", "거래품목", "결제유형", "거래처명", "지출금액", "적요"]
CATS = {"서비스제공": ["데크수리", "카오디오수리", "앰프수리", "탈부착", "내비게이션수리"],
        "상품판매": ["카오디오", "앰프", "스피커", "블랙박스"], "기타": ["체크입금", "카드"]}
MAKERS = ["벤츠", "BMW", "아우디", "폭스바겐", "볼보", "레인지로버", "현대", "N/A"]
CLIENTS = ["개인", "방문", "동서카오디오(대구)", "재즈카오디오(광주)", "현대카오디오(서울)",
           "수원테크(수원)", "옥션", "G마켓", "시험상사"]
EXP_ITEMS = [("제품원가", "제품매입"), ("비용", "식대"), ("비용", "차량유지비"), ("비용", "통신비"),
             ("생활", "마트"), ("생활", "간식비"), ("인건비", "급여"), ("기타", "기타")]


def kpi_block(ws, title):
    ws["B2"], ws["B3"] = "시험용 경영관리 시스템", title
    for col, label in (("F", "전일 합계"), ("H", "전일대비"), ("J", "당월 합계"), ("L", "영업일수")):
        ws[f"{col}2"], ws[f"{chr(ord(col) + 1)}2"] = label, 0


def put_header(ws, head):
    for i, h in enumerate(head):
        ws.cell(row=5, column=2 + i, value=h)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--year", type=int, default=2025)
    ap.add_argument("--sales", type=int, default=6000)
    ap.add_argument("--expenses", type=int, default=4000)
    ap.add_argument("--mode", choices=["both", "monthly"], default="both")
    ap.add_argument("--seed", type=int, default=7)
    a = ap.parse_args()
    rng = random.Random(a.seed)
    y = a.year
    days = [date(y, 1, 1) + timedelta(days=i) for i in range((date(y + 1, 1, 1) - date(y, 1, 1)).days)]

    sales, exps = [], []
    for i in range(a.sales):
        d = rng.choice(days)
        t = rng.choice(list(CATS))
        pay = rng.choices(["입금", "카드", "현금"], weights=[5, 4, 1])[0]
        tax = "Y" if pay == "카드" or rng.random() < 0.2 else "N"
        amt = rng.randrange(1, 200) * 10000 + rng.choice([0, 0, 0, 517])
        vat = amt / 11 if (pay == "카드" or tax == "Y") else 0
        client = rng.choice(CLIENTS)
        recv = amt if (client not in ("개인", "방문") and pay == "입금" and rng.random() < 0.08) else None
        payable = rng.randrange(1, 5) * 10000 if rng.random() < 0.01 else None
        sales.append({"d": d, "type": t, "cat": rng.choice(CATS[t]), "maker": rng.choice(MAKERS),
                      "model": rng.choice(["", "FCIC", "IMC", "NBT"]), "client": client, "amt": amt,
                      "pay": pay, "tax": tax, "vat": vat, "acct": rng.choice(["법인", "기업", "카드"]),
                      "memo": f"작업{i % 997}", "recv": recv, "payable": payable})
    # 같은 날 같은 내용의 정당한 중복 거래 (모두 등록되어야 함)
    for s in rng.sample(sales, 25):
        sales.append(dict(s))
    for i in range(a.expenses):
        d = rng.choice(days)
        et, item = rng.choice(EXP_ITEMS)
        exps.append({"d": d, "type": et, "item": item, "pay": rng.choice(["법인체크", "농협체크", "기업체크", "현금", "카드"]),
                     "client": rng.choice(["해덕", "마트", "식당", "주유소", ""]),
                     "amt": rng.randrange(1, 80) * 1000 + rng.choice([0, 400]), "memo": f"지출{i % 499}"})
    sales.sort(key=lambda r: r["d"])
    exps.sort(key=lambda r: r["d"])

    # 특이 행: 문자 날짜·문자 금액(정상 인식 대상), 잘못된 날짜(건너뜀 대상), 금액 0 행
    sales[10]["d_text"] = sales[10]["d"].strftime("%Y.%m.%d")
    sales[11]["amt_text"] = f"{sales[11]['amt']:,}"
    bad_sale = {"d": None, "d_raw": f"{y}-02-30", "type": "서비스제공", "cat": "데크수리", "maker": "N/A",
                "model": "", "client": "개인", "amt": 50000, "pay": "입금", "tax": "N", "vat": 0,
                "acct": "법인", "memo": "잘못된 날짜", "recv": None, "payable": None}

    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Instruction")["B2"] = "시험용 통합문서 (실데이터 아님)"
    form = wb.create_sheet("매출입력")
    for cell, v in (("C7", "영업일자"), ("E7", datetime(y, 1, 2)), ("C9", "매출유형"), ("E9", "서비스제공"),
                    ("C15", "거래통화"), ("E15", "CNY"), ("G15", "CNY 기준환율"), ("I15", 168.03),
                    ("C19", "매출액"), ("E19", 336060)):
        form[cell] = v
    ana = wb.create_sheet(f"{y - 1}년 경영분석")  # 원본처럼 제목 연도와 데이터 연도가 다를 수 있음
    ana["B2"] = f"{y - 1}년 경영분석"
    ana["B4"], ana["D4"] = "연간 누적 매출", 0
    for col, h in (("B", "영업월"), ("D", "매출액"), ("F", "지출액"), ("H", "영업이익")):
        ana[f"{col}15"] = h
    for m in range(1, 13):
        sh = wb.create_sheet(f"매출분석_{m}월")
        for i, h in enumerate(["영업월", "영업일자", "요일", "매출액", "전일대비"]):
            sh.cell(row=5, column=2 + i, value=h)
        sh.cell(row=6, column=2, value=f"{m}월")
        sh.cell(row=6, column=3, value=datetime(y, m, 1))
        sh.cell(row=6, column=5, value=0)

    def sale_row(ws, r, s):
        d = s["d"]
        ws.cell(row=r, column=2, value=int(f"{d.year}{d.month}") if d else None)
        ws.cell(row=r, column=3, value=s.get("d_text") or (datetime(d.year, d.month, d.day) if d else s["d_raw"]))
        ws.cell(row=r, column=4, value=WD[d.weekday()] if d else "")
        vals = [s["type"], s["cat"], s["maker"], s["model"], s["client"], s.get("amt_text") or s["amt"],
                s["amt"] if s["pay"] != "카드" else None, s["amt"] if s["pay"] == "카드" else None,
                s["vat"], s["amt"] - s["vat"], s["acct"], s["pay"], s["tax"], s["memo"], s["recv"], s["payable"]]
        for i, v in enumerate(vals):
            if v not in (None, ""):
                ws.cell(row=r, column=5 + i, value=v)

    def exp_row(ws, r, e):
        d = e["d"]
        ws.cell(row=r, column=2, value=int(f"{d.year}{d.month}"))
        ws.cell(row=r, column=3, value=datetime(d.year, d.month, d.day))
        ws.cell(row=r, column=4, value=WD[d.weekday()])
        for i, v in enumerate([e["type"], e["item"], e["pay"], e["client"], e["amt"], e["memo"]]):
            if v not in (None, ""):
                ws.cell(row=r, column=5 + i, value=v)

    def fill(ws, rows, writer, blank_cols):
        r = 6
        for row in rows:
            writer(ws, r, row)
            r += 1
        for _ in range(15):  # 수식만 있는 빈 서식 행(금액 칸이 0)
            for c in blank_cols:
                ws.cell(row=r, column=c, value=0)
            r += 1
        return r

    by_month_s, by_month_e = defaultdict(list), defaultdict(list)
    for s in sales:
        by_month_s[s["d"].month].append(s)
    for e in exps:
        by_month_e[e["d"].month].append(e)

    main_s = wb.create_sheet("매출집계")
    kpi_block(main_s, "매출관리")
    put_header(main_s, SALES_HEAD)
    if a.mode == "both":
        fill(main_s, sales, sale_row, (10, 13, 14))
    for m in range(1, 13):
        ws = wb.create_sheet(f"매출집계_{m}월")
        put_header(ws, SALES_HEAD)
        rows = by_month_s[m] + ([bad_sale] if m == 2 else [])
        end = fill(ws, rows, sale_row, (10, 13, 14))
        if m == 3:  # 날짜 없는 합계 행 (건너뛰고 '확인할 점'으로 알려야 함)
            ws.cell(row=end + 1, column=9, value="합계")
            ws.cell(row=end + 1, column=10, value=sum(s["amt"] for s in by_month_s[m]))
    main_e = wb.create_sheet("지출집계")
    kpi_block(main_e, "지출관리")
    put_header(main_e, EXP_HEAD)
    if a.mode == "both":
        fill(main_e, exps, exp_row, ())
    for m in range(1, 13):
        ws = wb.create_sheet(f"지출집계_{m}월")
        put_header(ws, EXP_HEAD)
        fill(ws, by_month_e[m], exp_row, ())
    wb.create_sheet("지출입력")["C7"] = "영업일자"
    wb.create_sheet("코드관리")["C1"] = "매출유형"

    # 경영분석 월별 합계 = 정상 내역의 월 합계 (엑셀 수식이 계산해 둔 값처럼)
    exp_months = {}
    for m in range(1, 13):
        inc = sum(s["amt"] for s in by_month_s[m])
        ex = sum(e["amt"] for e in by_month_e[m])
        ana.cell(row=15 + m, column=2, value=m)
        ana.cell(row=15 + m, column=4, value=inc)
        ana.cell(row=15 + m, column=6, value=ex)
        ana.cell(row=15 + m, column=8, value=inc - ex)
        exp_months[f"{y}-{m:02d}"] = {"income_count": len(by_month_s[m]), "income": inc,
                                       "expense_count": len(by_month_e[m]), "expense": ex,
                                       "receivable": sum(s["recv"] or 0 for s in by_month_s[m])}
    ana.cell(row=28, column=2, value="합계")
    ana.cell(row=28, column=4, value=sum(v["income"] for v in exp_months.values()))
    ana.cell(row=30, column=2, value="월 평균")

    wb.save(a.out)
    truth = {"mode": a.mode, "year": y, "sales": len(sales), "expenses": len(exps),
             "income_total": sum(s["amt"] for s in sales), "expense_total": sum(e["amt"] for e in exps),
             "receivable_total": sum(s["recv"] or 0 for s in sales),
             "payable_note_rows": sum(1 for s in sales if s["payable"]),
             "months": exp_months}
    with open(a.out.rsplit(".", 1)[0] + ".json", "w", encoding="utf-8") as f:
        json.dump(truth, f, ensure_ascii=False, indent=1)
    print(json.dumps({k: v for k, v in truth.items() if k != "months"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
